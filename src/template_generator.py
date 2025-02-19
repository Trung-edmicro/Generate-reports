from docx import Document
from fpdf import FPDF

def docx_to_fpdf_code(docx_path):
    doc = Document(docx_path)
    pdf_code = []
    
    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        
        if para.style.name.startswith("Heading"):
            pdf_code.append("pdf.set_font('DejaVu', 'B', 14)")
            pdf_code.append(f'pdf.cell(0, 8, "{text}", ln=True, align="C")')
            pdf_code.append("pdf.ln(5)")
            pdf_code.append("pdf.set_font('DejaVu', size=11)")

        match para.alignment:
            case None:
                if any(run.bold for run in para.runs):
                    if len(text) > 80:
                        pdf_code.append("pdf.set_font('DejaVu', 'B')")
                        pdf_code.append(f'pdf.multi_cell(0, 8, "{text}")')
                        pdf_code.append("pdf.ln(2)")
                        pdf_code.append("pdf.set_font('DejaVu', size=11)")
                    else:        
                        pdf_code.append("pdf.set_font('DejaVu', 'B')")
                        pdf_code.append(f'pdf.cell(0, 8, "{text}", ln=True)')
                        pdf_code.append("pdf.ln(2)")
                        pdf_code.append("pdf.set_font('DejaVu', size=11)")
                elif any(run.italic for run in para.runs):
                    if len(text) > 80:
                        pdf_code.append("pdf.set_font('DejaVu', 'I')")
                        pdf_code.append(f'pdf.multi_cell(0, 8, "{text}")')
                        pdf_code.append("pdf.ln(2)")
                        pdf_code.append("pdf.set_font('DejaVu', size=11)")
                    else:        
                        pdf_code.append("pdf.set_font('DejaVu', 'I')")
                        pdf_code.append(f'pdf.cell(0, 8, "{text}", ln=True)')
                        pdf_code.append("pdf.ln(2)")
                        pdf_code.append("pdf.set_font('DejaVu', size=11)")
                elif len(text) > 80:
                    pdf_code.append(f'pdf.multi_cell(0, 8, "{text}")')
                    pdf_code.append("pdf.ln(2)")
                else:
                    pdf_code.append(f'pdf.cell(0, 8, "{text}")')
                    pdf_code.append("pdf.ln(2)")
            case 1:
                if any(run.bold for run in para.runs):
                    pdf_code.append("pdf.set_font('DejaVu', 'B')")
                    pdf_code.append(f'pdf.cell(0, 8, "{text}", ln=True, align="C")')
                    pdf_code.append("pdf.ln(2)")
                    pdf_code.append("pdf.set_font('DejaVu', size=11)")
                elif any(run.italic for run in para.runs):
                    pdf_code.append("pdf.set_font('DejaVu', 'I')")
                    pdf_code.append(f'pdf.cell(0, 8, "{text}", ln=True, align="C")')
                    pdf_code.append("pdf.ln(2)")
                    pdf_code.append("pdf.set_font('DejaVu', size=11)")
                else:
                    pdf_code.append(f'pdf.cell(0, 8, "{text}", ln=True, align="C")')
                    pdf_code.append("pdf.ln(2)")
            case 2:
                if any(run.bold for run in para.runs):
                    pdf_code.append("pdf.set_font('DejaVu', 'B')")
                    pdf_code.append(f'pdf.cell(0, 8, "{text}", ln=True, align="R")')
                    pdf_code.append("pdf.ln(2)")
                    pdf_code.append("pdf.set_font('DejaVu', size=11)")
                elif any(run.italic for run in para.runs):
                    pdf_code.append("pdf.set_font('DejaVu', 'I')")
                    pdf_code.append(f'pdf.cell(0, 8, "{text}", ln=True, align="R")')
                    pdf_code.append("pdf.ln(2)")
                    pdf_code.append("pdf.set_font('DejaVu', size=11)")
                else:
                    pdf_code.append(f'pdf.cell(0, 8, "{text}", ln=True, align="R")')
                    pdf_code.append("pdf.ln(2)")
            case 3:
                return "Căn đều hai bên"
            case _:
                return f"Chưa có định dạng của {text}"
    
    pdf_code.append("pdf.output('output.pdf')")
    
    return '\n'.join(pdf_code)

# Chạy hàm với file template.docx
code_result = docx_to_fpdf_code("data/input/template_1.docx")
print(code_result)
